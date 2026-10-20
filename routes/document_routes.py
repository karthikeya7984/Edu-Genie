import io
import json
from typing import List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from sqlalchemy.orm import Session
import pypdf
import docx

import models
import schemas
from database import get_db
from utils.security import get_current_user
from services.ai_service import EduGenieAI
from routes.ai_routes import get_user_ai_params

router = APIRouter(prefix="/api/documents", tags=["Document Modules"])

MAX_FILE_SIZE = 15 * 1024 * 1024  # 15 MB

def extract_text_from_file(filename: str, file_bytes: bytes) -> str:
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    extracted_text = ""
    
    if ext == "pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
            extracted_text = "\n\n".join(text_parts)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read PDF file: {str(e)}"
            )
            
    elif ext in ("docx", "doc"):
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            text_parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text_parts.append(row_text)
            extracted_text = "\n\n".join(text_parts)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to read Word document: {str(e)}"
            )
            
    elif ext in ("txt", "md", "csv", "json"):
        try:
            extracted_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                extracted_text = file_bytes.decode("latin-1")
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to decode text file: {str(e)}"
                )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, TXT, or MD file."
        )

    clean_text = extracted_text.strip()
    if len(clean_text) < 15:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded file appears to be empty or does not contain readable text."
        )
    return clean_text


@router.post("/upload", response_model=schemas.DocumentDetailResponse)
async def upload_document(
    file: UploadFile = File(...),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    # Read file contents
    file_bytes = await file.read()
    file_size = len(file_bytes)
    
    if file_size > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum allowed size of 15MB. (File size: {file_size / (1024*1024):.1f}MB)"
        )
        
    filename = file.filename or "Uploaded_Document.pdf"
    file_type = filename.split(".")[-1].lower() if "." in filename else "unknown"
    
    # Extract text
    text_content = extract_text_from_file(filename, file_bytes)
    char_count = len(text_content)
    
    # Store in database
    new_doc = models.Document(
        user_id=current_user.id,
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        char_count=char_count,
        extracted_text=text_content
    )
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)
    
    # Generate 384-D Vector Embeddings for dense RAG retrieval
    try:
        from services.rag_service import RAGService
        RAGService.index_document(db, new_doc.id, text_content)
    except Exception as e:
        print(f"Vector embedding indexing error: {e}")
        
    return new_doc


@router.get("", response_model=List[schemas.DocumentResponse])
def get_user_documents(
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    docs = db.query(models.Document)\
        .filter(models.Document.user_id == current_user.id)\
        .order_by(models.Document.created_at.desc())\
        .all()
    return docs


@router.get("/{doc_id}", response_model=schemas.DocumentDetailResponse)
def get_document(
    doc_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return doc


@router.delete("/{doc_id}")
def delete_document(
    doc_id: int,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    
    db.delete(doc)
    db.commit()
    return {"status": "success", "message": f"Document '{doc.filename}' deleted successfully."}


@router.post("/{doc_id}/summarize", response_model=schemas.AIResponse)
def summarize_document(
    doc_id: int,
    req: schemas.DocumentSummarizeRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    params = get_user_ai_params(db, current_user.id)
    summary_text = EduGenieAI.document_summarize(
        document_text=doc.extracted_text,
        summary_type=req.summary_type or "Comprehensive",
        temperature=params["temperature"],
        model_name=params["model"]
    )
    
    # Cache summary in document record
    doc.summary = summary_text
    db.commit()
    
    # Save in chat history as well
    history = models.ChatHistory(
        user_id=current_user.id,
        module="document_summarize",
        prompt=f"Summarize document: {doc.filename} ({req.summary_type})",
        response=summary_text
    )
    db.add(history)
    db.commit()
    
    return {"module": "document_summarize", "data": summary_text}
    

@router.post("/{doc_id}/stream-summarize")
def stream_summarize_document(
    doc_id: int,
    req: schemas.DocumentSummarizeRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    params = get_user_ai_params(db, current_user.id)
    return StreamingResponse(
        EduGenieAI.stream_document_summarize(
            document_text=doc.extracted_text,
            summary_type=req.summary_type or "Comprehensive",
            temperature=params["temperature"],
            model_name=params["model"],
            doc_id=doc.id,
            user_id=current_user.id
        ),
        media_type="text/event-stream"
    )


@router.post("/{doc_id}/qa", response_model=schemas.AIResponse)
def ask_document(
    doc_id: int,
    req: schemas.DocumentQARequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    if not req.question or not req.question.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Question cannot be empty.")
        
    params = get_user_ai_params(db, current_user.id)
    answer = EduGenieAI.document_qa(
        document_text=doc.extracted_text,
        question=req.question.strip(),
        temperature=params["temperature"],
        model_name=params["model"],
        doc_id=doc.id,
        db=db
    )
    
    # Save in chat history
    history = models.ChatHistory(
        user_id=current_user.id,
        module="document_qa",
        prompt=f"Doc ({doc.filename}): {req.question}",
        response=answer
    )
    db.add(history)
    db.commit()
    
    return {"module": "document_qa", "data": answer}

from fastapi.responses import StreamingResponse

@router.post("/{doc_id}/stream-qa")
def stream_ask_document(
    doc_id: int,
    req: schemas.DocumentQARequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    if not req.question or not req.question.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Question cannot be empty.")
        
    params = get_user_ai_params(db, current_user.id)
    return StreamingResponse(
        EduGenieAI.stream_document_qa(
            document_text=doc.extracted_text,
            question=req.question.strip(),
            temperature=params["temperature"],
            model_name=params["model"],
            doc_id=doc.id,
            db=db,
            user_id=current_user.id
        ),
        media_type="text/event-stream"
    )


@router.post("/{doc_id}/quiz")
def quiz_from_document(
    doc_id: int,
    req: schemas.DocumentQuizRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    params = get_user_ai_params(db, current_user.id)
    questions_json = EduGenieAI.document_quiz(
        document_text=doc.extracted_text,
        difficulty=req.difficulty or "Medium",
        quiz_type=req.quiz_type or "MCQs",
        num_questions=req.num_questions or 5,
        temperature=params["temperature"],
        model_name=params["model"]
    )
    
    # Save the generated quiz to the database
    new_quiz = models.Quiz(
        topic=f"Document: {doc.filename}",
        difficulty=req.difficulty or "Medium",
        quiz_type=req.quiz_type or "MCQs",
        questions_json=questions_json
    )
    db.add(new_quiz)
    db.commit()
    db.refresh(new_quiz)
    
    try:
        parsed_questions = json.loads(questions_json)
    except Exception:
        parsed_questions = []
        
    return {
        "status": "success",
        "quiz_id": new_quiz.id,
        "topic": new_quiz.topic,
        "difficulty": new_quiz.difficulty,
        "quiz_type": new_quiz.quiz_type,
        "questions": parsed_questions
    }


@router.post("/{doc_id}/flashcards")
def flashcards_from_document(
    doc_id: int,
    req: schemas.DocumentFlashcardsRequest,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    doc = db.query(models.Document)\
        .filter(models.Document.id == doc_id, models.Document.user_id == current_user.id)\
        .first()
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        
    params = get_user_ai_params(db, current_user.id)
    flashcards_json = EduGenieAI.document_flashcards(
        document_text=doc.extracted_text,
        num_cards=req.num_cards or 5,
        temperature=params["temperature"],
        model_name=params["model"]
    )
    
    try:
        parsed_cards = json.loads(flashcards_json)
    except Exception:
        parsed_cards = []
        
    return {
        "status": "success",
        "filename": doc.filename,
        "flashcards": parsed_cards
    }
